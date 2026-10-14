"""Raspberry Pi Pico センサー入門スライドを生成するスクリプト。

使い方:
    pip install python-pptx
    python make_slides.py

出力:pico_sensor_intro.pptx(このファイルと同じフォルダ)
内容を直したいときは、下の「スライドの内容」の部分を編集して再実行する。
"""
import pathlib
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).parent
OUT = HERE / "pico_sensor_intro.pptx"

# ---------------------------------------------------------------- デザイン設定
FONT = "Meiryo"            # 本文
CODE_FONT = "Consolas"     # コード:ソースコード用の等幅フォント(英数字)
CODE_FONT_JA = "BIZ UDGothic"  # コード中の日本語(コメントなど)。等幅で Windows 10 以降に標準で入っている
BODY_PT = 24
CODE_PT = 18
CODE_LINE_PT = 17          # コード1行の高さ(枠の位置計算に使う)
SHORT_CODE_PT = 22         # 短いコード(数行の例)は大きく見せる
CODE_MAX_LINES = 17        # 1スライドに載せるコードの最大行数

C_TEXT = RGBColor(0x22, 0x22, 0x22)
C_SUB = RGBColor(0x66, 0x66, 0x66)
C_ACCENT = RGBColor(0x1F, 0x6F, 0xB2)
C_CODE_BG = RGBColor(0xF4, 0xF6, 0xF8)
C_COMMENT = RGBColor(0x2E, 0x7D, 0x32)
C_WARN_BG = RGBColor(0xFF, 0xF4, 0xE0)
C_WARN = RGBColor(0xD9, 0x6C, 0x00)
C_TIP_BG = RGBColor(0xE8, 0xF1, 0xFB)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_LED = {"赤": RGBColor(0xE5, 0x39, 0x35), "黄": RGBColor(0xF9, 0xA8, 0x25),
         "緑": RGBColor(0x43, 0xA0, 0x47)}

# 制御構造ごとの強調色
KW_COLORS = {
    "while": RGBColor(0x8E, 0x24, 0xAA),
    "if": RGBColor(0xE5, 0x39, 0x35),
    "for": RGBColor(0x1E, 0x88, 0xE5),
    "def": RGBColor(0x43, 0xA0, 0x47),
}

SW, SH = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_TOP = Inches(1.35)
CONTENT_W = SW - 2 * MARGIN

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- 部品(ヘルパー)
def set_font(run, size=BODY_PT, bold=False, color=C_TEXT, name=FONT):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, name
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    ja = CODE_FONT_JA if name == CODE_FONT else name   # 日本語部分のフォント
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", ja)


def add_rich(paragraph, text, size=BODY_PT, color=C_TEXT, name=FONT, bold=False):
    """**太字** と `コード` を解釈して run を追加する。"""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        r = paragraph.add_run()
        if part.startswith("**"):
            r.text = part[2:-2]
            set_font(r, size, True, C_ACCENT if color == C_TEXT else color, name)
        elif part.startswith("`"):
            r.text = part[1:-1]
            set_font(r, size, bold, RGBColor(0xC6, 0x28, 0x28), CODE_FONT)
        else:
            r.text = part
            set_font(r, size, bold, color, name)


def textbox(slide, x, y, w, h, lines, size=BODY_PT, color=C_TEXT, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, bullet=False, space=6):
    x, y, w, h = (Emu(int(v)) for v in (x, y, w, h))   # 座標は整数に
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        level = 0
        if isinstance(line, tuple):
            line, level = line
        prefix = ""
        if bullet:
            prefix = "・" if level == 0 else "　　– "
        add_rich(p, prefix + line, size - 2 * level, color)
    return tb


def bullets(slide, items, x=MARGIN, y=CONTENT_TOP, w=CONTENT_W, h=Inches(5.6), size=BODY_PT):
    return textbox(slide, x, y, w, h, items, size=size, bullet=True, space=10)


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1.5):
    x, y, w, h = (Emu(int(v)) for v in (x, y, w, h))   # 座標は整数に
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    return s


def label_box(slide, x, y, w, h, text, fill=C_TIP_BG, line=C_ACCENT, size=20, color=C_TEXT,
              bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = box(slide, x, y, w, h, fill, line, shape)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.1))
    for i, t in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        add_rich(p, t, size, color, bold=bold)
    return s


def callout(slide, text, y, h=Inches(0.8), kind="warn", x=MARGIN, w=CONTENT_W, size=22):
    fill, line = (C_WARN_BG, C_WARN) if kind == "warn" else (C_TIP_BG, C_ACCENT)
    icon = "⚠ " if kind == "warn" else "★ "
    s = box(slide, x, y, w, h, fill, line, MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    for i, t in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        add_rich(p, (icon if i == 0 else "") + t, size)
    return s


def line(slide, x1, y1, x2, y2, color=C_TEXT, width=2.5):
    # 座標は必ず整数にする(小数が入ると PowerPoint でファイルが開けなくなる)
    x1, y1, x2, y2 = (Emu(int(v)) for v in (x1, y1, x2, y2))
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def table(slide, rows, x, y, w, col_w=None, size=20, row_h=Inches(0.55), header=True):
    nr, nc = len(rows), len(rows[0])
    t = slide.shapes.add_table(nr, nc, x, y, w, row_h * nr).table
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            t.columns[i].width = Emu(int(w * cw / total))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            is_head = header and r == 0
            add_rich(p, str(val), size, C_WHITE if is_head else C_TEXT, bold=is_head)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_ACCENT if is_head else (
                RGBColor(0xFF, 0xFF, 0xFF) if r % 2 else RGBColor(0xF2, 0xF6, 0xFB))
    return t


def calc(slide, title, formula, numbers, answer, y, x=MARGIN, w=CONTENT_W, size=22):
    """計算を「式 → 数値を入れる → 答え」の3段で見せる。"""
    if title:
        textbox(slide, x, y, w, Inches(0.5), [f"**{title}**"], size=size)
        y += Inches(0.55)
    steps = [("式", formula, C_TIP_BG), ("数値を入れる", numbers, C_TIP_BG),
             ("答え", answer, RGBColor(0xFF, 0xF8, 0xD6))]
    h = Inches(0.62)
    for i, (lab, txt, fill) in enumerate(steps):
        yy = y + i * (h + Inches(0.12))
        label_box(slide, x, yy, Inches(2.2), h, lab, C_ACCENT, C_ACCENT, size - 2, C_WHITE, True)
        s = box(slide, x + Inches(2.3), yy, w - Inches(2.3), h, fill, None)
        tf = s.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.2)
        add_rich(tf.paragraphs[0], txt, size, bold=(lab == "答え"))
    return y + 3 * (h + Inches(0.12))


def text_width_in(s, size=CODE_PT):
    """等幅フォントでの文字列の幅(インチ)をおおまかに見積もる。"""
    units = sum(1.0 if ord(ch) > 0x2E80 else 0.55 for ch in s)   # Consolas の英数字は 0.55 文字幅
    return units * size / 72


def code_block(slide, lines, x=MARGIN, y=CONTENT_TOP, w=CONTENT_W, size=CODE_PT):
    """コードを表示し、ブロックの下端の y を返す。size で文字の大きさを変えられる。"""
    line_pt = int(size * 1.25)   # 1行の高さ(14pt → 17pt)
    h = Inches(0.25) + Pt(line_pt) * len(lines)
    bg = box(slide, x, y, w, h, C_CODE_BG, RGBColor(0xD0, 0xD7, 0xDE), line_w=1)
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    pad = Inches(0.15)
    tf.margin_left = tf.margin_right = pad
    tf.margin_top = tf.margin_bottom = Inches(0.12)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(line_pt)
        p.alignment = PP_ALIGN.LEFT
        # コードとコメントを分けて色を付ける
        m = re.match(r"^(.*?)(#.*)?$", ln)
        code_part, comment = m.group(1), m.group(2) or ""
        if code_part:
            r = p.add_run()
            r.text = code_part
            set_font(r, size, False, C_TEXT, CODE_FONT)
        if comment:
            r = p.add_run()
            r.text = comment
            set_font(r, size, False, C_COMMENT, CODE_FONT)
        if not ln:
            r = p.add_run()
            r.text = " "
            set_font(r, size, name=CODE_FONT)
        # 制御構造の行を色付き枠で囲む
        kw = re.match(r"^\s*(while|if|elif|else|for|def)\b", code_part)
        if kw:
            key = {"elif": "if", "else": "if"}.get(kw.group(1), kw.group(1))
            indent = len(code_part) - len(code_part.lstrip())
            fx = x + pad + Inches(text_width_in(" " * indent, size)) - Inches(0.05)
            fw = Inches(text_width_in(code_part.rstrip(), size) - text_width_in(" " * indent, size) + 0.1)
            fy = y + Inches(0.12) + Pt(line_pt) * i - Pt(1)
            f = box(slide, fx, fy, fw, Pt(line_pt + 1), None, KW_COLORS[key],
                    MSO_SHAPE.ROUNDED_RECTANGLE, line_w=2)
    return y + h


def kw_legend(slide, used):
    """左下に制御構造の色の凡例を出す。"""
    if not used:
        return
    x = MARGIN
    textbox(slide, x, SH - Inches(0.5), Inches(1.4), Inches(0.4), ["色付き枠:"], size=12, color=C_SUB,
            anchor=MSO_ANCHOR.MIDDLE)
    x += Inches(1.3)
    for key in ["while", "if", "for", "def"]:
        if key not in used:
            continue
        label_box(slide, x, SH - Inches(0.48), Inches(0.9), Inches(0.36), key, C_WHITE,
                  KW_COLORS[key], 12, KW_COLORS[key], True)
        x += Inches(1.0)


slide_no = 0


def new_slide(title, notes, part=None):
    global slide_no
    slide_no += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = C_WHITE
    textbox(s, MARGIN, Inches(0.25), CONTENT_W, Inches(0.8), [title], size=32 if len(title) <= 22 else 28,
            anchor=MSO_ANCHOR.MIDDLE)
    s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    box(s, MARGIN, Inches(1.08), CONTENT_W, Pt(3), C_ACCENT)
    footer = f"{part}  |  {slide_no}" if part else str(slide_no)
    textbox(s, MARGIN, SH - Inches(0.45), CONTENT_W, Inches(0.35), [footer], size=12,
            color=C_SUB, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def code_slides(title, code, notes_list, part):
    """長いコードを空行の位置で分割して、複数スライドに載せる(コメントは削らない)。"""
    lines = code.strip("\n").split("\n")
    # 1スライドに最大 CODE_MAX_LINES 行まで詰める。区切りは、なるべく後ろのほうの空行にする
    chunks = []
    rest = lines
    while rest:
        while rest and not rest[0].strip():          # 先頭の空行は飛ばす
            rest = rest[1:]
        if len(rest) <= CODE_MAX_LINES:
            chunks.append(rest)
            break
        window = rest[:CODE_MAX_LINES + 1]
        # 区切る場所の優先順位:①後ろのほうの空行 ②後ろのほうのコメント行の手前 ③ちょうど最大行数
        near = CODE_MAX_LINES - 4
        blanks = [i for i in range(near, CODE_MAX_LINES + 1) if not window[i].strip()]
        comments = [i for i in range(near, CODE_MAX_LINES) if window[i].strip().startswith("#")]
        cut = (blanks or comments or [CODE_MAX_LINES])[-1]
        chunks.append(rest[:cut])
        rest = rest[cut:]
    chunks = [c for c in chunks if c]
    # 最後のスライドが数行だけにならないよう、最後の2枚を空行の位置で半分ずつに分け直す
    if len(chunks) > 1 and len(chunks[-1]) < 6:
        both = chunks[-2] + [""] + chunks[-1]
        blanks = [i for i, l in enumerate(both) if not l.strip()
                  and i <= CODE_MAX_LINES and len(both) - i - 1 <= CODE_MAX_LINES]
        if blanks:
            cut = min(blanks, key=lambda i: abs(i - len(both) / 2))
            chunks[-2:] = [both[:cut], both[cut + 1:]]
    n = len(chunks)
    for k, ch in enumerate(chunks):
        t = title if n == 1 else f"{title}({k + 1}/{n})"
        s = new_slide(t, notes_list[min(k, len(notes_list) - 1)], part)
        used = {({"elif": "if", "else": "if"}.get(m.group(1), m.group(1)))
                for l in ch for m in [re.match(r"^\s*(while|if|elif|else|for|def)\b", l)] if m}
        kw_legend(s, used)
        # 長いコードは文字が大きいので、枠を左右に広げて1行が折り返さないようにする
        code_block(s, ch, x=Inches(0.3), w=SW - Inches(0.6))
        if k < n - 1:
            textbox(s, MARGIN, SH - Inches(0.45), CONTENT_W - Inches(2.8), Inches(0.35),
                    ["→ 次のスライドに続く"], size=12, color=C_SUB, align=PP_ALIGN.RIGHT)


def wiring(slide, pairs, part_name, x=MARGIN, y=CONTENT_TOP + Inches(0.1), w=Inches(6.6),
           row=Inches(0.62), colors=None):
    """Pico と部品を線でつなぐ簡単な接続図。pairs = [(Pico 側, 部品側), ...]"""
    n = len(pairs)
    h = row * n + Inches(0.3)
    pw = Inches(2.6)
    label_box(slide, x, y - Inches(0.45), pw, Inches(0.4), "Pico", None, None, 18, C_SUB, True)
    label_box(slide, x + w - pw, y - Inches(0.45), pw, Inches(0.4), part_name, None, None, 18,
              C_SUB, True)
    box(slide, x, y, pw, h, RGBColor(0xE8, 0xF5, 0xE9), RGBColor(0x2E, 0x7D, 0x32),
        MSO_SHAPE.ROUNDED_RECTANGLE)
    box(slide, x + w - pw, y, pw, h, RGBColor(0xF3, 0xF3, 0xF3), C_SUB, MSO_SHAPE.ROUNDED_RECTANGLE)
    for i, (a, b) in enumerate(pairs):
        yy = y + Inches(0.15) + row * i
        col = (colors or {}).get(i, C_TEXT)
        textbox(slide, x, yy, pw, row, [a], size=18, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, x + w - pw, yy, pw, row, [b], size=18, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        line(slide, x + pw, yy + row / 2, x + w - pw, yy + row / 2, col, 3)


PIN_NOTE = "配布したピン配置図で、GP番号と物理ピン番号の両方を確かめよう"

# ================================================================ スライドの内容
P1, P2, P3, P4, P5, P6, P7, P8 = ("第1部 はじめに", "第2部 制御構造", "第3部 LED と信号機",
                                  "第4部 タクトスイッチ", "第5部 人感センサー", "第6部 超音波センサー",
                                  "第7部 OLED(発展)", "第8部 まとめ")

# ---- S1 タイトル
s = prs.slides.add_slide(BLANK)
slide_no += 1
s.background.fill.solid()
s.background.fill.fore_color.rgb = C_WHITE
box(s, 0, Inches(2.2), SW, Inches(2.6), C_ACCENT)
textbox(s, MARGIN, Inches(2.4), CONTENT_W, Inches(1.3), ["Raspberry Pi Pico でセンサーを動かそう"],
        size=40, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
textbox(s, MARGIN, Inches(3.7), CONTENT_W, Inches(0.8), ["〜信号機・人感センサー・超音波センサー〜"],
        size=28, color=C_WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i, (name, col) in enumerate(C_LED.items()):
    box(s, Inches(5.67) + Inches(0.75) * i, Inches(5.4), Inches(0.5), Inches(0.5), col, None,
        MSO_SHAPE.OVAL)
s.notes_slide.notes_text_frame.text = (
    "みなさん、こんにちは。今日は Raspberry Pi Pico という小さなコンピュータを使って、"
    "信号機を光らせたり、人や物との距離をセンサーで測ったりします。"
    "プログラムで現実の物が動くと、とても楽しいですよ。最後まで一緒にがんばりましょう。")

# ---- S2 今日のゴール
s = new_slide("今日のゴール", (
    "今日のゴールは3つです。1つ目は、Pico で LED を光らせたりセンサーの値を読んだりすること。"
    "2つ目は、while・if・for・関数を使って、プログラムの動きの流れを作れるようになること。"
    "3つ目は、抵抗やプルアップといった電気の基本が、なぜ必要なのかを説明できるようになることです。"), P1)
for i, (num, txt) in enumerate([
        ("1", "Pico で **LED を光らせる**/**センサーの値を読む**"),
        ("2", "`while` / `if` / `for` / 関数 を使って「動きの流れ」を作れる"),
        ("3", "電気の基本(**抵抗・電圧・プルアップ**)がなぜ必要か説明できる")]):
    yy = CONTENT_TOP + Inches(0.3) + Inches(1.6) * i
    label_box(s, MARGIN, yy, Inches(1.0), Inches(1.0), num, C_ACCENT, C_ACCENT, 36, C_WHITE, True,
              MSO_SHAPE.OVAL)
    textbox(s, MARGIN + Inches(1.3), yy, CONTENT_W - Inches(1.3), Inches(1.0), [txt],
            anchor=MSO_ANCHOR.MIDDLE, size=26)

# ---- S3 Pico と GPIO
s = new_slide("Raspberry Pi Pico と GPIO", (
    "Pico はマイコン、つまり小さなコンピュータです。GPIO は Pico が電気を出したり読み取ったりできる出入口です。"
    "3.3V が HIGH で 1、0V が LOW で 0 です。"
    "ここで一番大事なのは、GP番号と物理ピン番号は別物だということ。物理3番は GND で、GP3 ではありません。"
    "つまずきやすい点:配線ミスの多くはこの取り違えです。配布したピン配置図を必ず手元に置いてください。"), P1)
bullets(s, [
    "**マイコン** = 小さなコンピュータ",
    "**GPIO** =「Pico が電気を ON/OFF したり、読み取ったりできる出入口」",
    ("出力(OUT):電気を出す 例)LED を光らせる", 1),
    ("入力(IN):電気を読み取る 例)ボタンを調べる", 1),
    "**3.3V が HIGH(1)**、**0V が LOW(0)**",
], h=Inches(3.4))
table(s, [["番号の種類", "意味", "例"],
          ["GP番号", "プログラムで使う番号", "`Pin(3, Pin.OUT)` → GP3"],
          ["物理ピン番号", "基板上の位置(1〜40)", "GP3 は物理5番"]],
      MARGIN, Inches(4.4), CONTENT_W, [2, 3.5, 3.5], size=20, row_h=Inches(0.5))
callout(s, "物理3番は GND であって GP3 ではない!\n" + PIN_NOTE, Inches(6.15), Inches(0.85), size=18)

# ---- S4 Thonny
s = new_slide("Thonny の準備", (
    "Thonny を開いたら、まず右下のインタプリタを MicroPython (Raspberry Pi Pico) にしてください。"
    "緑の三角が実行、赤い四角が停止です。print の結果は下の Shell に出てきます。"
    "つまずきやすい点:Pico が認識されないときは、USB ケーブルが充電専用でないか確認しましょう。"), P1)
bullets(s, [
    "Pico を USB でパソコンにつなぐ",
    "右下のインタプリタを **「MicroPython (Raspberry Pi Pico)」** にする",
    "Shell に `>>>` が出ればつながっている",
], h=Inches(2.2))
table(s, [["ボタン", "意味"],
          ["▶ 実行(緑の三角)", "プログラムを Pico で動かす"],
          ["■ 停止(赤い四角)", "動いているプログラムを止める"],
          ["Shell(画面下)", "`print` の結果が表示される"]],
      MARGIN, Inches(3.7), CONTENT_W, [1, 1.4], size=22, row_h=Inches(0.6))
callout(s, "Pico が見つからない → USB ケーブルが「充電専用」でないか確認", Inches(6.3),
        Inches(0.6), size=20)

step1 = (HERE.parent / "01-python-basics" / "step1_on.py").read_text(encoding="utf-8")
step2 = (HERE.parent / "01-python-basics" / "step2_on_off.py").read_text(encoding="utf-8")
step3 = (HERE.parent / "01-python-basics" / "step3_blink.py").read_text(encoding="utf-8")
# ---- まずは LED をつけてみよう(3ステップ)
for title, code, pts, note in [
    ("まずは LED をつけてみよう ①つける", step1, [
        "`Pin(\"LED\", Pin.OUT)`:本体の LED を **出力** にする", "`led.value(1)`:**1 で点灯**",
        "やってみよう:`1` を `0` にすると?(→ 消える)"],
     "まずは3行だけのプログラムで、Pico 本体の LED をつけてみましょう。"
     "Pin(\"LED\", Pin.OUT) で本体の LED を出力にして、value(1) で点灯です。"
     "つまずきやすい点:右下のインタプリタが MicroPython (Raspberry Pi Pico) になっていないと動きません。"),
    ("まずは LED をつけてみよう ②つけて消す", step2, [
        "`time.sleep(1)`:**1秒待つ**", "上から順に:点灯 → 待つ → 消灯",
        "やってみよう:`1` を `3` にすると?(→ 3秒光る)"],
     "次は、つけて1秒待って消します。time.sleep(1) が1秒待つ命令です。"
     "プログラムは上から順番に実行されるので、点灯、待つ、消灯の順に動きます。"),
    ("まずは LED をつけてみよう ③点滅", step3, [
        "`while True:` の中(字下げした4行)が **ずっとくり返される**", "止めるときは **停止ボタン(■)**",
        "やってみよう:`0.5` を `0.1` にすると?(→ 速く点滅)"],
     "最後は while True で、つけて消してをずっとくり返します。これで点滅になります。"
     "つまずきやすい点:while の中の行は字下げ(スペース4つ)をそろえましょう。止めるときは停止ボタンです。"),
]:
    s = new_slide(title, note, P1)
    lines = code.strip("\n").split("\n")
    bottom = code_block(s, lines, y=CONTENT_TOP + Inches(0.1), size=20)
    kw_legend(s, {"while"} if "while" in code else set())
    bullets(s, pts, y=bottom + Inches(0.25), h=Inches(2), size=22)

# ---- S5 順次
s = new_slide("プログラムは上から順に動く(順次)", (
    "プログラムは書いた順番に、上から1行ずつ実行されます。これを順次といいます。"
    "当たり前のようですが、順番を入れかえると動きが変わる、ということを意識しておきましょう。"), P2)
code_block(s, ["print(\"1行目\")", "print(\"2行目\")   # 1行目のあとに実行される",
               "print(\"3行目\")   # 2行目のあとに実行される"], y=CONTENT_TOP + Inches(0.3),
           w=Inches(8.3), size=SHORT_CODE_PT)
for i, t in enumerate(["1行目", "2行目", "3行目"]):
    yy = CONTENT_TOP + Inches(0.2) + Inches(1.2) * i
    label_box(s, Inches(9.7), yy, Inches(2.9), Inches(0.8), t, size=24)
    if i < 2:
        line(s, Inches(11.15), yy + Inches(0.8), Inches(11.15), yy + Inches(1.2), C_ACCENT, 3)
bullets(s, ["**1行ずつ、上から順番に** 実行される", "順番を入れかえると、動きも変わる"],
        y=Inches(5.0), h=Inches(1.5))

# ---- S6 while True
s = new_slide("while True: 無限ループ", (
    "マイコンは電源が入っている間、ずっと動き続けるのが基本です。そのために while True を使います。"
    "True は「ずっと正しい」という意味なので、ループは終わりません。止めるときは Thonny の停止ボタンを押します。"), P2)
kw_legend(s, {"while"})
code_block(s, ["while True:          # True なので終わらない", "    print(\"ずっと表示\")",
               "    time.sleep(1)    # 1秒待つ"], y=CONTENT_TOP + Inches(0.3), w=Inches(8.5), size=SHORT_CODE_PT)
bullets(s, ["マイコンは **「ずっと動き続ける」** のが基本",
            "`while True:` の中は、終わらずにくり返される",
            "止めるには Thonny の **停止ボタン(■)**"], y=Inches(3.45), h=Inches(2.5))

# ---- S7 if
s = new_slide("if 条件分岐", (
    "if は「もし〜なら」という意味です。条件によって処理を分けたいときに使います。"
    "このあと、ボタンが押されていたら、距離が10cm未満なら、という形で何度も出てきます。"), P2)
kw_legend(s, {"if"})
code_block(s, ["if distance < 10:      # もし距離が10cm未満なら", "    print(\"近い!\")",
               "else:                  # そうでなければ", "    print(\"遠い\")"],
           y=CONTENT_TOP + Inches(0.3), w=Inches(8.5), size=SHORT_CODE_PT)
bullets(s, ["「**もし** ボタンが押されていたら」", "「**もし** 距離が10cm未満なら」",
            "条件によって、する処理を分ける"], y=Inches(3.85), h=Inches(2.5))

# ---- S8 for
s = new_slide("for 繰り返し", (
    "for は、決まった回数や決まったものを順番に処理するときに使います。"
    "信号機では、赤・黄・緑の LED を1個ずつ取り出して点灯させるのに使います。"), P2)
kw_legend(s, {"for"})
code_block(s, ["for i in range(3):          # i に 0, 1, 2 が順に入る", "    print(i)",
               "", "for p in (red, yellow, green):   # LED を1個ずつ取り出す", "    p.value(1)"],
           y=CONTENT_TOP + Inches(0.3), w=Inches(10), size=SHORT_CODE_PT)
bullets(s, ["**決まった回数** をくり返す", "**決まったもの** を順に処理する(LED を1個ずつ点灯)"],
        y=Inches(4.25), h=Inches(2))

# ---- S9 関数
s = new_slide("関数 def", (
    "よく使う処理には、def で名前を付けておけます。これが関数です。"
    "今日は全部の LED を消す all_off や、距離を測る read_cm が出てきます。"
    "return を使うと、関数の結果を呼び出した側に返せます。"), P2)
kw_legend(s, {"def"})
code_block(s, ["def double(x):", "    return x * 2       # x を2倍した結果を返す", "",
               "print(double(5))       # 10 と表示される"], y=CONTENT_TOP + Inches(0.3),
           w=Inches(9), size=SHORT_CODE_PT)
bullets(s, ["よく使う処理に **名前を付ける**", ("例:`all_off()`(全部消す)、`read_cm()`(距離を測る)", 1),
            "`return` で **結果を返す**"], y=Inches(3.85), h=Inches(2.5))

# ---- S10 sleep
s = new_slide("time.sleep() 待つ命令", (
    "sleep は待つ命令です。sleep(1) で1秒、sleep_ms は100分の1ではなく1000分の1秒単位、sleep_us は100万分の1秒単位です。"
    "そして大事なのが、待っている間は Pico は他のことが何もできない、ということ。"
    "これがあとでボタンを使うときに問題になります。覚えておいてください。"), P2)
table(s, [["書き方", "待つ時間"], ["`time.sleep(1)`", "1秒"],
          ["`time.sleep_ms(100)`", "100ミリ秒 = 0.1秒"],
          ["`time.sleep_us(10)`", "10マイクロ秒 = 0.00001秒"]],
      MARGIN, CONTENT_TOP + Inches(0.2), CONTENT_W, [1, 1], size=24, row_h=Inches(0.7))
callout(s, "待っている間は **他のことができない!**\n→ あとのボタン・割り込みで、この問題が出てくるよ",
        Inches(4.6), Inches(1.3), kind="tip", size=24)

# ---- S11 LED と抵抗
s = new_slide("LED は抵抗なしでつないではいけない", (
    "LED は電流が流れすぎると壊れてしまいます。最悪の場合、Pico のほうも傷みます。"
    "そこで、電流を抑えるために電流制限抵抗を LED と直列、つまり一列につなぎます。"
    "LED を単体で光らせるときは、220Ω くらいの抵抗が必要だと覚えてください。"), P3)
bullets(s, ["LED は **電流が流れすぎると壊れる**(最悪の場合 Pico 側も傷む)",
            "電流を抑えるために **電流制限抵抗** を **直列**(一列)に入れる",
            "単体 LED を光らせるときは **220Ω 程度の抵抗が必要**"], h=Inches(2.8))
yy = Inches(5.0)
label_box(s, Inches(1.2), yy, Inches(2.2), Inches(0.9), "GP2\n(3.3V)", RGBColor(0xE8, 0xF5, 0xE9),
          RGBColor(0x2E, 0x7D, 0x32), 18)
label_box(s, Inches(4.4), yy + Inches(0.15), Inches(2.0), Inches(0.6), "抵抗 220Ω", C_WHITE, C_TEXT, 18,
          shape=MSO_SHAPE.RECTANGLE)
label_box(s, Inches(7.4), yy, Inches(1.4), Inches(0.9), "LED", C_LED["赤"], C_LED["赤"], 20, C_WHITE,
          True, MSO_SHAPE.OVAL)
label_box(s, Inches(9.8), yy, Inches(2.2), Inches(0.9), "GND", RGBColor(0xE8, 0xF5, 0xE9),
          RGBColor(0x2E, 0x7D, 0x32), 18)
for a, b in [(3.4, 4.4), (6.4, 7.4), (8.8, 9.8)]:
    line(s, Inches(a), yy + Inches(0.45), Inches(b), yy + Inches(0.45), C_TEXT, 3)

# ---- S12 オームの法則
s = new_slide("抵抗値の計算(オームの法則)", (
    "オームの法則は、電圧イコール電流かける抵抗です。LED は光るときに約2V を使います。これを順方向電圧といいます。"
    "Pico の3.3V から LED の2V を引いた1.3V が抵抗にかかります。それを220Ωで割ると約0.006A、つまり6mA です。"
    "LED の目安の20mA より十分小さいので安全です。抵抗がないと、電流を止めるものがなくなってしまいます。"), P3)
textbox(s, MARGIN, CONTENT_TOP - Inches(0.1), CONTENT_W, Inches(0.9), [
    "**オームの法則:V = I × R → I = V ÷ R**",
    "LED は光るとき約 **2V** を使う(順方向電圧 Vf・赤LEDの目安)"], size=20, space=2)
y2 = calc(s, "① 抵抗にかかる電圧", "電源 − LED の電圧", "3.3V − 2.0V", "1.3V",
          CONTENT_TOP + Inches(0.85), w=Inches(5.9), size=20)
calc(s, "② 流れる電流", "I = V ÷ R", "1.3V ÷ 220Ω", "≒ 0.006A = 6mA",
     CONTENT_TOP + Inches(0.85), x=Inches(6.8), w=Inches(5.9), size=20)
callout(s, "6mA は安全な範囲(LED は目安 20mA 以下、Pico の GPIO は控えめに)。\n"
        "抵抗なしだと? → 電流を止めるものがなく、大きな電流が流れてしまう!", Inches(5.4), Inches(1.1),
        size=20)

# ---- S13 信号機モジュール
s = new_slide("今日の信号機モジュールは抵抗入り", (
    "今日使う信号機モジュールには、基板に抵抗が組み込まれています。なので GPIO に直接つないで大丈夫です。"
    "でも、部品を単体で使うときは220Ωを入れる、ということは必ず覚えておいてください。"
    "つまずきやすい点:GND は物理3番です。GP3 と取り違えないように、ピン配置図で確認しましょう。"), P3)
wiring(s, [("GP2(物理4)", "R(赤)"), ("GP3(物理5)", "Y(黄)"), ("GP4(物理6)", "G(緑)"),
           ("GND(物理3)", "GND")], "信号機モジュール",
       colors={0: C_LED["赤"], 1: C_LED["黄"], 2: C_LED["緑"]}, y=CONTENT_TOP + Inches(0.6))
bullets(s, ["基板に **抵抗が組み込まれている** → GPIO に **直結 OK**",
            "部品を単体で使うときは **220Ω を入れる**", "GND は **物理3番**(GP3 ではない!)"],
        x=Inches(7.6), y=CONTENT_TOP + Inches(0.2), w=Inches(5.2), h=Inches(3.5), size=22)
callout(s, PIN_NOTE, Inches(5.9), Inches(0.6), kind="tip", size=20)

# ---- S14 信号機コード
TRAFFIC = '''
# ---- 必要な道具(モジュール)を読み込む ----
from machine import Pin   # Pin:Pico のピンを操作するための道具
import time               # time:待ち時間を作るための道具

# ---- ピンの設定 ----
# Pin(番号, Pin.OUT) で「このピンは電気を出す(出力)」と決める
red    = Pin(2, Pin.OUT)   # GP2 に赤のLED
yellow = Pin(3, Pin.OUT)   # GP3 に黄のLED
green  = Pin(4, Pin.OUT)   # GP4 に緑のLED

# ---- 関数:全部のLEDを消す ----
def all_off():
    # 赤・黄・緑を1つずつ取り出して p に入れ、順番に消す
    for p in (red, yellow, green):
        p.value(0)         # 0 = 電気を止める = 消灯

# ---- 配線チェック:1個ずつ点灯 ----
# (名前, ピン) の組を順番に取り出して処理する
for name, p in (("赤", red), ("黄", yellow), ("緑", green)):
    all_off()              # いったん全部消す
    print(name, "点灯")    # Shell に今どの色かを表示
    p.value(1)             # 1 = 電気を流す = 点灯
    time.sleep(1)          # 1秒待つ
all_off()                  # チェックが終わったら全部消す

# ---- 信号のサイクル(ずっと繰り返す) ----
while True:                # True なので終わらない(無限ループ)
    green.value(1)         # 緑を点ける
    time.sleep(5)          # 5秒待つ
    green.value(0)         # 緑を消す

    yellow.value(1)        # 黄を点ける
    time.sleep(2)          # 2秒待つ
    yellow.value(0)        # 黄を消す

    red.value(1)           # 赤を点ける
    time.sleep(5)          # 5秒待つ
    red.value(0)           # 赤を消す → また緑に戻る
'''
code_slides("信号機コード(配線チェック+サイクル)", TRAFFIC, [
    "前半はピンの設定と、全部消す関数、そして配線チェックです。コメントを1行ずつ読んでみてください。"
    "for で赤・黄・緑を1個ずつ点けて、Shell に出た色と光った色が合っているか確認します。",
    "後半は while True で、緑5秒、黄2秒、赤5秒をずっと繰り返します。"
    "つまずきやすい点:色が合わないときは、GP番号と物理ピン番号の取り違えを疑いましょう。"], P3)

# ---- S15 読み解き
s = new_slide("信号機コードの読み解き", (
    "Pin(2, Pin.OUT) で GP2 を出力にしています。all_off は for を使って3つの LED を消す関数です。"
    "最初の for は配線チェック、while True が信号のサイクルです。"
    "では実習です。秒数を変えて、自分だけの信号機にしてみましょう。"
    "つまずきやすい点:time.sleep の数字は秒、sleep_ms はミリ秒なので単位に注意。"), P3)
table(s, [["コード", "意味"], ["`Pin(2, Pin.OUT)`", "GP2 を **出力** に設定"],
          ["`def all_off():` + `for`", "全部消す **関数**"],
          ["`for name, p in ...`", "1個ずつ点灯 → **配線と色が合っているか確認**"],
          ["`while True:`", "緑 → 黄 → 赤 を **くり返す**"]],
      MARGIN, CONTENT_TOP + Inches(0.1), CONTENT_W, [1.2, 1.6], size=22, row_h=Inches(0.7))
callout(s, "実習:秒数を変えてみよう!(例:緑3秒・黄1秒・赤3秒)", Inches(5.4), Inches(0.8),
        kind="tip", size=24)

# ---- S16 タクトスイッチ
s = new_slide("タクトスイッチのしくみ", (
    "タクトスイッチは、押している間だけ電気が通るスイッチです。片側を GP5、物理7番に、もう片側を GND、物理8番につなぎます。"
    "つまずきやすい点:タクトスイッチは足が4本あります。対角線の足を使うと確実です。"), P4)
bullets(s, ["**押している間だけ** 電気が通るスイッチ", "足は4本 → **対角線の足** を使うと確実"],
        h=Inches(1.6))
wiring(s, [("GP5(物理7)", "片側"), ("GND(物理8)", "もう片側")], "タクトスイッチ",
       y=Inches(3.6))
callout(s, PIN_NOTE, Inches(5.9), Inches(0.6), kind="tip", size=20)

# ---- S17 フローティング
s = new_slide("何もつながっていない入力は「ふらふら」する", (
    "ボタンを押していないとき、GP5 はどこにもつながっていません。これをフローティング、浮いている状態といいます。"
    "このとき値は HIGH か LOW か決まらず、指で触っただけで変わることもあります。"
    "線がアンテナのように、まわりの電気のノイズを拾ってしまうイメージです。"), P4)
bullets(s, ["押していないとき、GP5 は **どこにもつながっていない**(**フローティング**)",
            "HIGH か LOW か決まらない", "触っただけで値が変わることも"], w=Inches(7), h=Inches(3))
label_box(s, Inches(8.3), Inches(4.4), Inches(1.8), Inches(0.8), "GP5", RGBColor(0xE8, 0xF5, 0xE9),
          RGBColor(0x2E, 0x7D, 0x32), 22, bold=True)
line(s, Inches(9.2), Inches(4.4), Inches(9.2), Inches(2.9), C_TEXT, 3)
for i, t in enumerate(["〰 ノイズ", "〰 ノイズ", "〰 ノイズ"]):
    textbox(s, Inches(9.5) + Inches(0.6) * (i % 2), Inches(1.6) + Inches(0.9) * i, Inches(2.6),
            Inches(0.5), [t], size=20, color=C_WARN)
textbox(s, Inches(7.2), Inches(5.3), Inches(4.5), Inches(0.5), ["線がアンテナのようにノイズを拾う"],
        size=18, color=C_SUB, align=PP_ALIGN.CENTER)
callout(s, "値が「1? 0? 1? 0?」とふらふら…", Inches(5.9), Inches(0.6), size=20, w=Inches(6.4))

# ---- S18 プルアップ
s = new_slide("プルアップ抵抗とは", (
    "そこで、入力ピンを抵抗を通して3.3V に引っ張り上げておきます。これがプルアップです。"
    "押していないときは3.3V につながって1、押すと GND につながって0になります。押すと0、というのが感覚と逆なので注意です。"
    "Pico には内部にプルアップ抵抗があるので、Pin.PULL_UP と書くだけで使えます。抵抗があるので、押しても3.3V と GND がショートしません。"), P4)
code_block(s, [" 3.3V", "  │", " [抵抗]  ← Pico の内部にある", "  │", "  ├───── GP5(入力)", "  │",
               "  ┴ スイッチ", "  │", " GND"], x=MARGIN, y=CONTENT_TOP + Inches(0.1), w=Inches(4.8),
           size=18)
table(s, [["ボタン", "つながる先", "値"], ["押していない", "3.3V", "**HIGH(1)**"],
          ["押している", "GND", "**LOW(0)**"]],
      Inches(5.8), CONTENT_TOP + Inches(0.1), Inches(6.9), [1.3, 1, 1.1], size=22, row_h=Inches(0.6))
bullets(s, ["`Pin.PULL_UP` で内部のプルアップ抵抗を使える(外付け不要)",
            "抵抗があるので、押しても 3.3V と GND が **ショートしない**"],
        x=Inches(5.8), y=Inches(3.5), w=Inches(6.9), h=Inches(2), size=22)
callout(s, "**「押すと 0」** になる点が逆なので注意!", Inches(5.9), Inches(0.7), size=22)

# ---- S19 チャタリング
s = new_slide("チャタリング", (
    "スイッチは押した瞬間に、中の金属の接点が細かく ON と OFF を繰り返すことがあります。これをチャタリングといいます。"
    "1回押しただけなのに、何回も押されたように見えてしまいます。"
    "対策として、押されたら20ミリ秒待って、もう一度確認します。それでも押されていたら本当に押したと判断します。"), P4)
bullets(s, ["押した瞬間に接点が **細かく ON/OFF をくり返す**", "1回押したのに、何回も押したように見える"],
        h=Inches(1.6))
# 波形の図
x0, yb, yt = Inches(1.0), Inches(4.6), Inches(3.4)
pts = [(0, 1), (1.5, 1), (1.5, 0), (1.8, 0), (1.8, 1), (2.0, 1), (2.0, 0), (2.3, 0), (2.3, 1), (2.45, 1),
       (2.45, 0), (6.0, 0)]
for (a, va), (b, vb) in zip(pts, pts[1:]):
    line(s, x0 + Inches(a), yb if va == 0 else yt, x0 + Inches(b), yb if vb == 0 else yt, C_ACCENT, 3)
textbox(s, x0 + Inches(1.3), Inches(4.75), Inches(1.6), Inches(0.5), ["バタバタ!"], size=18,
        color=C_WARN)
textbox(s, Inches(0.2), Inches(3.25), Inches(0.8), Inches(0.4), ["1"], size=18)
textbox(s, Inches(0.2), Inches(4.45), Inches(0.8), Inches(0.4), ["0"], size=18)
label_box(s, Inches(7.8), Inches(3.3), Inches(4.9), Inches(1.6),
          "対策:**少し待って(20ms)**\n**もう一度確認する**", C_TIP_BG, C_ACCENT, 24)

# ---- S20 ポーリング版コード
POLLING = '''
# ---- 必要な道具を読み込む ----
from machine import Pin
import time

# ---- ピンの設定 ----
red    = Pin(2, Pin.OUT)   # GP2:赤(出力)
yellow = Pin(3, Pin.OUT)   # GP3:黄(出力)
green  = Pin(4, Pin.OUT)   # GP4:緑(出力)

# GP5 をボタン用の「入力(IN)」にする
# PULL_UP:内部のプルアップ抵抗を使う
#   押していない → 1(HIGH)
#   押している   → 0(LOW) ← 押すと 0 になる点に注意!
button = Pin(5, Pin.IN, Pin.PULL_UP)

# ---- 関数:全部のLEDを消す ----
def all_off():
    for p in (red, yellow, green):
        p.value(0)

# ---- スタート時の状態 ----
all_off()
green.value(1)             # 車用の信号は緑で待つ
print("待機中(緑)")

# ---- メインの繰り返し ----
while True:
    # ボタンの状態を読む。0 なら押されている
    if button.value() == 0:
        time.sleep_ms(20)              # 20ミリ秒待つ(チャタリング対策)
        if button.value() == 0:        # 20ミリ秒後もまだ押されていたら「本当に押した」と判断
            print("ボタン押下")
            time.sleep(3)              # 押してからすぐには変わらず、3秒待つ

            # 緑 → 黄
            green.value(0)
            yellow.value(1)
            print("黄")
            time.sleep(2)              # 黄は2秒

            # 黄 → 赤
            yellow.value(0)
            red.value(1)
            print("赤(歩行者OK)")
            time.sleep(5)              # 赤は5秒(この間に歩行者がわたる)

            # 赤 → 緑(元に戻る)
            red.value(0)
            green.value(1)
            print("緑に復帰")

            # ボタンが押しっぱなしなら、離されるまでここで待つ
            # (待たないと、すぐに次のサイクルが始まってしまう)
            while button.value() == 0:
                time.sleep_ms(10)

    # 10ミリ秒ごとにボタンを見に行く(これがポーリング)
    time.sleep_ms(10)
'''
code_slides("押しボタン信号(ポーリング版)コード", POLLING, [
    "ボタンのピン設定に注目してください。Pin.IN で入力、Pin.PULL_UP で内部プルアップを使います。押すと0になります。",
    "メインの繰り返しです。ボタンが0なら20ミリ秒待って、もう一度確認しています。これがチャタリング対策です。"
    "つまずきやすい点:プルアップなので、== 0 で押された判定になります。",
    "最後に、ボタンが離されるまで待っています。そして10ミリ秒ごとにボタンを見に行きます。これがポーリングです。"], P4)

# ---- S21 ポーリング版読み解き
s = new_slide("ポーリング版の読み解き", (
    "ポーリングとは、何度も見に行って確認する方式のことです。if の中に if が入っているのがチャタリング対策、"
    "while button.value() == 0 は離すまで待つ部分です。"
    "では実験です。黄や赤の最中にボタンを押してみてください。まったく反応しませんね。sleep している間はボタンを見ていないからです。"), P4)
bullets(s, ["**ポーリング** =「何度も見に行って確認する」方式",
            "`if` の中に `if`(**入れ子**)→ チャタリング対策",
            "`while button.value() == 0:` → **離すまで待つ**"], h=Inches(2.6))
callout(s, "実験:黄・赤の最中にボタンを押してみよう\n→ まったく反応しない! `sleep` 中は **見ていない** から",
        Inches(4.3), Inches(1.4), size=24)

# ---- S22 割り込み
s = new_slide("割り込み(irq)とは", (
    "割り込みは、何かが起きたら知らせてもらう方式です。ポーリングは郵便受けを何度も見に行くこと、"
    "割り込みはインターホンが鳴ったら出ること、とたとえられます。"
    "IRQ_FALLING は、HIGH から LOW に落ちた瞬間、つまり押した瞬間です。"
    "割り込みで呼ばれる関数をハンドラといい、フラグを立てるだけの短い処理にするのがコツです。"), P4)
label_box(s, MARGIN, CONTENT_TOP + Inches(0.2), Inches(5.8), Inches(1.5),
          "ポーリング\n郵便受けを **何度も見に行く**", RGBColor(0xF3, 0xF3, 0xF3), C_SUB, 24)
label_box(s, Inches(6.9), CONTENT_TOP + Inches(0.2), Inches(5.8), Inches(1.5),
          "割り込み\n**インターホンが鳴ったら** 出る", C_TIP_BG, C_ACCENT, 24)
bullets(s, ["割り込み =「何かが起きたら **知らせてもらう**」方式",
            "`Pin.IRQ_FALLING`:HIGH → LOW に変わった瞬間(= **押した瞬間**)",
            "割り込みで呼ばれる関数(**ハンドラ**)は短く:**フラグを立てるだけ**"],
        y=Inches(3.4), h=Inches(3))

# ---- S23 irq版コード
IRQ = '''
# ---- 必要な道具を読み込む ----
from machine import Pin
import time

# ---- ピンの設定(ポーリング版と同じ) ----
red    = Pin(2, Pin.OUT)
yellow = Pin(3, Pin.OUT)
green  = Pin(4, Pin.OUT)
button = Pin(5, Pin.IN, Pin.PULL_UP)   # 押すと 0 になる

# ---- フラグ(目印)の変数 ----
# ボタンが押されたら True にする。最初は「押されていない」= False
pressed = False

# ---- 割り込みで呼ばれる関数(ハンドラ) ----
# ボタンが押された瞬間に、Pico が自動でこの関数を呼んでくれる
def on_press(pin):
    global pressed     # 関数の外にある pressed を書き換えるための宣言
    pressed = True     # 「押されたよ」と目印を立てるだけ(短く書くのが大事)

# ---- 割り込みの登録 ----
# IRQ_FALLING:1 → 0 に変わった瞬間(=押した瞬間)に on_press を呼ぶ
button.irq(trigger=Pin.IRQ_FALLING, handler=on_press)

# ---- 関数:全部のLEDを消す ----
def all_off():
    for p in (red, yellow, green):
        p.value(0)

# ---- スタート時の状態 ----
all_off()
green.value(1)
print("待機中(緑)")

# ---- メインの繰り返し ----
while True:
    # ボタンを直接見に行かず、目印(pressed)を確認する
    if pressed:
        pressed = False            # 目印を下ろしておく
        print("ボタン押下")
        time.sleep_ms(3000)        # 3秒待つ

        # 緑 → 黄
        green.value(0)
        yellow.value(1)
        print("黄")
        time.sleep_ms(2000)        # 黄は2秒

        # 黄 → 赤
        yellow.value(0)
        red.value(1)
        print("赤(歩行者OK)")
        time.sleep_ms(5000)        # 赤は5秒

        # 赤 → 緑
        red.value(0)
        green.value(1)
        print("緑に復帰")

        # 黄・赤の間に押されていても、割り込みで目印が立っている
        # ここで False に戻すと「その押下は無視」
        # この行を消すと「予約」として次のサイクルが始まる
        pressed = False

    time.sleep_ms(50)              # 0.05秒ごとに目印を確認
'''
code_slides("押しボタン信号(irq版)コード", IRQ, [
    "pressed というフラグ、目印の変数を用意します。on_press がハンドラで、目印を立てるだけです。"
    "button.irq で、押した瞬間に on_press を呼んでね、と登録しています。",
    "メインの繰り返しでは、ボタンを直接見に行かず、目印を確認しています。目印が立っていたら信号を変えます。",
    "最後の pressed = False で、黄や赤の間に押されたぶんを無視しています。"
    "つまずきやすい点:global を書き忘れると、ボタンを押しても何も起きません。"], P4)

# ---- S24 irq比較
s = new_slide("irq版の読み解きとポーリング版との比較", (
    "global pressed は、関数の外の変数を書き換えるための宣言です。"
    "ポーリングと割り込みをくらべると、割り込みは sleep 中に押されてもフラグとして残る、というのが大きな違いです。"
    "発展として、最後の pressed = False を消すと、押した分が予約になって次のサイクルにつながります。試してみてください。"), P4)
bullets(s, ["`global pressed`:関数の外の変数を書き換えるための宣言"], h=Inches(0.8))
table(s, [["", "ポーリング", "割り込み(irq)"], ["しくみ", "何度も見に行く", "起きたら知らせてもらう"],
          ["sleep 中の押下", "気づかない", "**フラグとして残る**"],
          ["わかりやすさ", "わかりやすい", "少し難しい"]],
      MARGIN, CONTENT_TOP + Inches(0.9), CONTENT_W, [1, 1.3, 1.3], size=22, row_h=Inches(0.7))
callout(s, "発展:最後の `pressed = False` を消すと?\n→「**予約**」として次のサイクルにつながる", Inches(5.0),
        Inches(1.3), kind="tip", size=22)

# ---- S25 人感センサー
s = new_slide("人感センサーのしくみ", (
    "人感センサーは、人の体から出る赤外線、つまり熱の変化を感じ取ります。検知すると OUT が HIGH、いなくなると LOW になります。"
    "電源を入れてから数十秒は不安定なので、ウォームアップが必要です。また検知後しばらくは HIGH が続きます。これをホールド時間といいます。"
    "つまずきやすい点:VCC・OUT・GND の並びは製品によって違うことがあるので、基板の印字を必ず確認してください。"), P5)
bullets(s, ["人の体の **赤外線(熱)の変化** を感じる",
            "検知 → OUT が **HIGH**", "いなくなる → **LOW**",
            "電源を入れて数十秒は不安定", ("→ **ウォームアップ** が必要", 1),
            "検知後しばらく HIGH が続く", ("→ **ホールド時間**", 1)], w=Inches(6.6), h=Inches(4.2),
        size=22)
wiring(s, [("3V3(OUT)(物理36)", "VCC"), ("GND(物理38)", "GND"), ("GP15(物理20)", "OUT")],
       "SB412A", x=Inches(7.3), w=Inches(5.4), y=CONTENT_TOP + Inches(0.6))
callout(s, "基板の印字で **VCC / OUT / GND の並び** を必ず確認!" + "  " + PIN_NOTE, Inches(5.7),
        Inches(0.9), size=20)

# ---- S26 人感センサーコード
PIR = '''
# ---- 必要な道具を読み込む ----
from machine import Pin
import time

# ---- ピンの設定 ----
# GP15 を人感センサーの入力にする
# PULL_DOWN:何もつながっていないときは 0(LOW)に固定する
pir = Pin(15, Pin.IN, Pin.PULL_DOWN)

# Pico 本体に付いているLEDを出力にする(番号ではなく "LED" と名前で指定する)
led = Pin("LED", Pin.OUT)

# ---- ウォームアップ ----
# 電源を入れてすぐはセンサーが不安定なので30秒待つ
print("ウォームアップ中...")
time.sleep(30)
print("開始")

# ---- 前回の値を覚えておく変数 ----
prev = 0   # 最初は「検知なし」として始める

# ---- メインの繰り返し ----
while True:
    now = pir.value()          # 今のセンサーの値(1:検知、0:検知なし)

    # 前回と値が変わったときだけ表示する
    # (毎回表示すると Shell が文字だらけになるため)
    if now != prev:
        if now == 1:
            print("検知!")
        else:
            print("検知なし")
        led.value(now)         # 検知中は本体LEDを点ける
        prev = now             # 今の値を「前回の値」として覚える

    time.sleep_ms(100)         # 0.1秒ごとに確認
'''
code_slides("人感センサーのコード", PIR, [
    "GP15 を入力にして、PULL_DOWN を使います。本体の LED は「LED」という名前で指定します。最初に30秒のウォームアップをします。"
    "つまずきやすい点:30秒間は何も表示されませんが、壊れていないので待ちましょう。",
    "メインでは、前回の値 prev と今の値 now をくらべて、変わったときだけ表示します。"], P5)

# ---- S27 読み解き
s = new_slide("人感センサーコードの読み解き", (
    "PULL_DOWN はプルアップの逆で、何もないときに LOW に固定します。"
    "prev と now をくらべて、変化したときだけ表示する、という考え方はいろいろなセンサーで使えます。"
    "実習として、検知したら信号機を赤にしてみましょう。"), P5)
table(s, [["", "プルアップ", "プルダウン"], ["何もないとき", "3.3V に引っ張る → **1**", "GND に引っ張る → **0**"],
          ["書き方", "`Pin.PULL_UP`", "`Pin.PULL_DOWN`"]],
      MARGIN, CONTENT_TOP + Inches(0.1), CONTENT_W, [1, 1.4, 1.4], size=22, row_h=Inches(0.65))
bullets(s, ["`prev`(前回の値)と `now`(今の値)をくらべて、**変化したときだけ** 表示する"],
        y=Inches(3.6), h=Inches(1.2))
callout(s, "実習:検知したら信号機を赤にしてみよう!", Inches(5.2), Inches(0.8), kind="tip", size=24)

# ---- S28 超音波のしくみ
s = new_slide("超音波センサーのしくみ", (
    "超音波センサーは、人には聞こえない40kHz の高い音を出して、物に当たって返ってくるまでの時間を測ります。"
    "コウモリやイルカと同じ考え方です。Trig は音を出す合図、Echo は音が返ってくるまでの時間だけ HIGH になります。"), P6)
bullets(s, ["人に聞こえない高い音(**40kHz**)を出し、**返ってくるまでの時間** を測る",
            "コウモリやイルカと同じ考え方",
            "**Trig**:音を出す合図/**Echo**:返ってくるまでの時間だけ HIGH"], h=Inches(2.6))
yy = Inches(4.5)
label_box(s, Inches(1.0), yy, Inches(2.4), Inches(1.4), "センサー\nHC-SR04", C_TIP_BG, C_ACCENT, 20)
label_box(s, Inches(10.0), yy, Inches(1.8), Inches(1.4), "物体", RGBColor(0xF3, 0xF3, 0xF3), C_SUB, 22,
          shape=MSO_SHAPE.RECTANGLE)
for dy, txt, col in [(0.35, "(音)→", C_ACCENT), (1.05, "←(反射)", C_WARN)]:
    line(s, Inches(3.5), yy + Inches(dy), Inches(9.9), yy + Inches(dy), col, 3)
    textbox(s, Inches(5.5), yy + Inches(dy) - Inches(0.45), Inches(2.5), Inches(0.4), [txt], size=18,
            color=col, align=PP_ALIGN.CENTER)

# ---- S29 距離の計算
s = new_slide("音の速さから距離を求める", (
    "音速は約340m/秒です。正確には331.5足す0.6かける気温で、気温によって変わります。"
    "単位をそろえると、340m/秒は34000cm/秒、さらに100万で割って0.034cm/μs です。"
    "距離は速さかける時間ですが、音は行って帰ってくるので2で割ります。0.034÷2 は0.017、これは58.8で割るのと同じです。"
    "だからコードでは t / 58 と書いています。たとえば580μs なら、580÷58で10cm です。"), P6)
textbox(s, MARGIN, CONTENT_TOP - Inches(0.05), CONTENT_W, Inches(0.5), [
    "音速 約 **340m/秒**(気温15℃付近。正確には **331.5 + 0.6 × 気温**)μs = 100万分の1秒"], size=20)
calc(s, "① 単位をそろえる", "m/秒 → cm/秒 → cm/μs", "340m/秒 = 34000cm/秒", "0.034cm/μs",
     CONTENT_TOP + Inches(0.5), w=Inches(6.1), size=18)
calc(s, "② 往復なので 2 で割る", "距離 = 時間 × 0.034 ÷ 2", "時間 × 0.017(= ÷ 58.8)",
     "距離 ≒ 時間 ÷ 58(`t / 58`)", CONTENT_TOP + Inches(0.5), x=Inches(6.9), w=Inches(5.8), size=18)
calc(s, "③ 例:Echo が 580μs", "距離 = 時間 ÷ 58", "580 ÷ 58", "10cm", Inches(4.6), size=18)

# ---- S30 5V と 3.3V
s = new_slide("5V と 3.3V の問題", (
    "HC-SR04 は5V で動くセンサーで、Echo からも5V の信号が出てきます。"
    "でも Pico の入力は3.3V までです。5V をそのまま入れると壊れるおそれがあります。"
    "そこで、電圧を下げる分圧回路を使います。"), P6)
label_box(s, Inches(1.0), Inches(2.2), Inches(3.4), Inches(1.6), "HC-SR04\nEcho から **5V**",
          RGBColor(0xFD, 0xEC, 0xEA), C_LED["赤"], 24)
label_box(s, Inches(8.9), Inches(2.2), Inches(3.4), Inches(1.6), "Pico\n入力は **3.3V まで**",
          RGBColor(0xE8, 0xF5, 0xE9), RGBColor(0x2E, 0x7D, 0x32), 24)
line(s, Inches(4.4), Inches(3.0), Inches(8.9), Inches(3.0), C_LED["赤"], 4)
textbox(s, Inches(5.0), Inches(2.3), Inches(3.4), Inches(0.6), ["5V そのままは ✕"], size=22,
        color=C_LED["赤"], align=PP_ALIGN.CENTER)
callout(s, "5V をそのまま入れると **壊れるおそれ**\n→ 電圧を下げる **「分圧回路」** を使う", Inches(4.6),
        Inches(1.4), size=26)

# ---- S31 分圧
s = new_slide("抵抗で電圧を分ける(分圧)", (
    "抵抗を2つ直列につなぐと、電圧は抵抗の大きさに比例して分かれます。"
    "公式は、出力電圧イコール入力電圧かける R2 割る R1足す R2 です。"
    "1kΩ と2kΩ なら、5V かける3分の2で約3.3V になります。水の流れを2つの坂に分けると、途中の高さは坂の長さで決まる、というイメージです。"
    "Trig は Pico から3.3V を出すので直結で OK。3.3V 対応品の HC-SR04P なら分圧は不要です。"
    "つまずきやすい点:1kΩ と2kΩ の位置を逆にしないこと。"), P6)
textbox(s, MARGIN, CONTENT_TOP - Inches(0.05), Inches(7.6), Inches(0.5),
        ["**出力電圧 = 入力電圧 × R2 ÷(R1 + R2)**"], size=22)
calc(s, None, "5V × R2 ÷(R1 + R2)", "5V × 2kΩ ÷(1kΩ + 2kΩ)= 5 × 2/3", "≒ 3.3V",
     CONTENT_TOP + Inches(0.55), w=Inches(7.6), size=20)
code_block(s, ["Echo ──[1kΩ]──┬── GP11", "              │", "            [2kΩ]", "              │",
               "             GND"], x=Inches(8.6), y=CONTENT_TOP, w=Inches(4.1), size=18)
textbox(s, Inches(8.6), Inches(3.2), Inches(4.1), Inches(0.9),
        ["R1 = 1kΩ(Echo 側)", "R2 = 2kΩ(GND 側)"], size=18, color=C_SUB)
bullets(s, ["たとえ:水の流れを2つの坂に分けると、途中の高さは坂の長さで決まる",
            "**Trig** は 3.3V の信号で動くので **直結 OK**",
            "3.3V 対応品(**HC-SR04P** など)なら分圧は不要"], y=Inches(4.15), h=Inches(2.6), size=22)

# ---- S32 配線まとめ
s = new_slide("超音波センサーの配線まとめ", (
    "配線をまとめます。VCC は VBUS の40番、GND は38番、Trig は GP10 の14番、Echo は分圧回路を通して GP11 の15番です。"
    "つまずきやすい点:VCC は3V3ではなく VBUS、5V につなぎます。Echo は必ず分圧回路を通しましょう。"), P6)
wiring(s, [("VBUS 5V(物理40)", "VCC"), ("GND(物理38)", "GND"), ("GP10(物理14)", "Trig"),
           ("GP11(物理15)", "Echo")], "HC-SR04", colors={3: C_WARN}, y=CONTENT_TOP + Inches(0.6),
       w=Inches(7.2))
label_box(s, Inches(3.55), CONTENT_TOP + Inches(0.75) + Inches(0.62) * 3, Inches(1.9), Inches(0.42),
          "分圧回路", C_WARN_BG, C_WARN, 16)
bullets(s, ["VCC は **VBUS(5V)**", "Echo は **分圧回路経由**", "Trig は直結 OK"], x=Inches(8.3),
        y=CONTENT_TOP + Inches(0.4), w=Inches(4.4), h=Inches(3), size=22)
callout(s, PIN_NOTE, Inches(5.8), Inches(0.6), kind="tip", size=20)

# ---- S33 超音波コード
ULTRA = '''
# ---- 必要な道具を読み込む ----
# time_pulse_us:ピンが HIGH になっている時間を μs(100万分の1秒)単位で測る道具
from machine import Pin, time_pulse_us
import time

# ---- ピンの設定 ----
trig = Pin(10, Pin.OUT)   # GP10:Trig(音を出す合図を送る → 出力)
echo = Pin(11, Pin.IN)    # GP11:Echo(返ってきた時間を受け取る → 入力)
                          # ※ Echo は分圧回路(1kΩ・2kΩ)を通してつなぐ

# ---- 関数:距離を測って cm で返す ----
def read_cm():
    # 合図の前に、いったん 0 にして落ち着かせる
    trig.value(0)
    time.sleep_us(2)

    # 10μs だけ 1 にする → センサーが超音波を出す
    trig.value(1)
    time.sleep_us(10)
    trig.value(0)

    # Echo が 1 になっている時間(音が往復した時間)を測る
    # 30000μs(=30ミリ秒)待っても返ってこなければあきらめる
    t = time_pulse_us(echo, 1, 30000)

    # 測れなかったとき、t はマイナスの値になる
    if t < 0:
        return None        # None =「値なし」を返す

    # 時間(μs)を距離(cm)に変える
    # 音速 0.034cm/μs × 時間 ÷ 2(往復なので)≒ 時間 ÷ 58
    return t / 58

# ---- メインの繰り返し ----
while True:
    d = read_cm()          # 距離を測る

    if d is None:
        print("測定不可")
    else:
        # {:.1f} は「小数第1位まで表示する」という意味
        print("{:.1f} cm".format(d))

    time.sleep_ms(200)     # 0.2秒ごとに測る
'''
code_slides("超音波センサーのコード", ULTRA, [
    "距離を測る処理を read_cm という関数にまとめています。Trig を10μs だけ1にして、超音波を出させます。",
    "time_pulse_us で Echo が HIGH の時間を測り、t / 58 で cm に変えて返します。メインでは0.2秒ごとに距離を表示します。"
    "つまずきやすい点:測定不可ばかり出るときは、配線と分圧回路を確認しましょう。"], P6)

# ---- S34 読み解き
s = new_slide("超音波センサーコードの読み解き", (
    "Trig を10μs だけ HIGH にするのが、音を出してという合図です。time_pulse_us は Echo が HIGH の時間を測り、30ms 以上はあきらめます。"
    "測れなかったときは None、値なしを返します。return t / 58 はさっき計算した式です。"
    "測定範囲は2cm から400cm くらいで、近すぎる物、斜めの面、やわらかい物は苦手です。"), P6)
table(s, [["コード", "意味"], ["Trig を 10μs だけ HIGH", "「音を出して!」の合図"],
          ["`time_pulse_us(echo, 1, 30000)`", "Echo が HIGH の時間を μs で測る。30ms 以上はあきらめる"],
          ["`if t < 0: return None`", "測れなかったときは「値なし」を返す"],
          ["`return t / 58`", "時間 → 距離(cm)の計算"]],
      MARGIN, CONTENT_TOP + Inches(0.1), CONTENT_W, [1.3, 1.7], size=20, row_h=Inches(0.7))
callout(s, "測定範囲は約 **2〜400cm**。近すぎる・斜めの面・やわらかい物は苦手", Inches(5.3), Inches(0.8),
        size=22)

# ---- S34-1 OLED と I2C
s = new_slide("OLED ディスプレイと I2C", (
    "ここからは発展です。OLED は1つ1つの点が自分で光る画面で、今回のものは横128、縦64ドットです。"
    "左上が0,0で、右に x、下に y が増えます。標準の文字は8×8ドットなので、1行16文字、8行まで書けます。"
    "OLED とは I2C という、SDA と SCL の2本の線でやりとりします。SSD1306 のアドレスはふつう0x3C、10進数で60です。"), P7)
bullets(s, ["OLED:**1つ1つの点が自分で光る** 画面(横 **128** × 縦 **64** ドット)",
            "左上が **(0, 0)**。右へ x、下へ y が増える",
            "文字は **8×8 ドット** → 1行 **16文字**・**8行** まで(英数字のみ)",
            "**I2C**:**SDA**(データ)と **SCL**(合図)の **2本の線** でやりとり",
            "SSD1306 のアドレスはふつう **0x3C(= 60)**"], h=Inches(3.6))
code_block(s, ["(0,0) ──────▶ x (0〜127)", "  │", "  ▼", "  y (0〜63)"], y=Inches(5.0),
           w=Inches(5.6), size=18)

# ---- S34-2 ライブラリのインストール
s = new_slide("ライブラリを入れよう(Thonny)", (
    "OLED を動かすには ssd1306 というライブラリが必要です。Thonny のツール、パッケージを管理から入れます。"
    "Pico の中に lib/ssd1306.py ができていれば OK です。"
    "つまずきやすい点:Pico をつないで、右下が MicroPython になっている状態でインストールしましょう。"), P7)
for i, txt in enumerate(["メニューの **ツール → パッケージを管理…** を開く",
                         "`ssd1306` で検索し、**micropython-ssd1306** を **インストール**",
                         "Pico の中に `lib/ssd1306.py` ができていれば OK"]):
    yy = CONTENT_TOP + Inches(0.3) + Inches(1.3) * i
    label_box(s, MARGIN, yy, Inches(0.9), Inches(0.9), str(i + 1), C_ACCENT, C_ACCENT, 28, C_WHITE, True,
              MSO_SHAPE.OVAL)
    textbox(s, MARGIN + Inches(1.2), yy, CONTENT_W - Inches(1.2), Inches(0.9), [txt], size=24,
            anchor=MSO_ANCHOR.MIDDLE)
callout(s, "Pico をつないだ状態(右下が MicroPython)でインストールしよう", Inches(5.6), Inches(0.7),
        kind="tip", size=22)

# ---- S34-3 配線
s = new_slide("OLED の配線", (
    "配線です。GND は23番、VCC は3V3の36番、SCL は GP19 の25番、SDA は GP18 の24番です。"
    "信号機やボタンとピンが重ならないように、I2C1 の GP18 と GP19 を使います。"
    "つまずきやすい点:端子の並びは製品によって違うので、OLED に書いてある文字を必ず見ましょう。"), P7)
wiring(s, [("GND(物理23)", "GND"), ("3V3(物理36)", "VCC"), ("GP19(物理25)", "SCL"),
           ("GP18(物理24)", "SDA")], "SSD1306", y=CONTENT_TOP + Inches(0.6), w=Inches(7.2))
bullets(s, ["VCC は **3V3** につなぐ", "**I2C1**(GP18・GP19)を使う", "並びは **製品ごとにちがう**"],
        x=Inches(8.3), y=CONTENT_TOP + Inches(0.4), w=Inches(4.4), h=Inches(3), size=22)
callout(s, PIN_NOTE, Inches(5.8), Inches(0.6), kind="tip", size=20)

# ---- S34-4 OLED コード(08-oled/oled.py を読み込む)
OLED = (HERE.parent / "08-oled" / "oled.py").read_text(encoding="utf-8")
code_slides("OLED のコード", OLED, [
    "I2C1 を GP18 と GP19 で使う準備をして、i2c.scan() でつながっている部品を調べます。60と出れば OK です。"
    "text で文字、rect で四角を描き、最後に show で画面に送ります。",
    "while True の中で、消す、描く、show をくり返して数字を書きかえます。"
    "つまずきやすい点:show を忘れると何も表示されません。"], P7)

# ---- S34-5 読み解き
s = new_slide("OLED コードの読み解き", (
    "大事なのは、text や rect は下書きに描くだけで、show を呼んだときにはじめて画面に出ることです。"
    "表示を書きかえるときは fill(0) で消してから描きます。数は str で文字に変えてから text に渡します。"
    "scan の結果が空っぽなら、SDA と SCL が逆になっていないか確認しましょう。"), P7)
table(s, [["コード", "意味"], ["`i2c.scan()`", "つながっている部品のアドレス。**[60]** なら OK"],
          ["`oled.fill(0)`", "画面を全部消す(0 = 黒、1 = 白)"],
          ["`oled.text(\"文字\", x, y)`", "文字を描く(下書き)"],
          ["`oled.show()`", "下書きを **画面に送る**"],
          ["`str(n)`", "数を文字に変える"]],
      MARGIN, CONTENT_TOP + Inches(0.1), CONTENT_W, [1.3, 1.7], size=20, row_h=Inches(0.65))
callout(s, "`[]` と出たら配線を確認(**SDA と SCL が逆** になっていない?)", Inches(5.6), Inches(0.8),
        size=22)

# ---- S35 まとめ
s = new_slide("今日のまとめ", (
    "今日学んだことをふり返りましょう。制御構造は順次・while・if・for・関数。"
    "電気では、LED には電流制限抵抗、プルアップとプルダウン、分圧。"
    "センサーは、人感センサーが赤外線、超音波センサーが音速かける時間割る2。ボタンはポーリングと割り込みでした。"), P8)
for i, (head, body) in enumerate([
        ("制御構造", "順次・`while`・`if`・`for`・関数"),
        ("電気の基本", "LED には電流制限抵抗(220Ω程度)/プルアップ・プルダウン/分圧"),
        ("センサー", "人感(赤外線)/超音波(音速 × 時間 ÷ 2)"),
        ("ボタン", "ポーリングと割り込み")]):
    yy = CONTENT_TOP + Inches(0.2) + Inches(1.3) * i
    label_box(s, MARGIN, yy, Inches(2.6), Inches(1.0), head, C_ACCENT, C_ACCENT, 24, C_WHITE, True)
    textbox(s, MARGIN + Inches(2.9), yy, CONTENT_W - Inches(2.9), Inches(1.0), [body], size=24,
            anchor=MSO_ANCHOR.MIDDLE)

# ---- S36 チャレンジ
s = new_slide("チャレンジ課題", (
    "最後にチャレンジ課題です。今日使った部品とプログラムを組み合わせてみましょう。"
    "ピン設定や関数は、それぞれのコードからコピーしてきて OK です。"
    "つまずきやすい点:部品が増えると GND が足りなくなります。ブレッドボードで GND をまとめましょう。"), P8)
for i, txt in enumerate([
        "超音波で **10cm 以内** に近づいたら信号機を **赤** にする",
        "人感センサーが反応したら **黄を点滅** させる",
        "ボタンを押したら超音波で距離を測り、**距離に応じて信号の色を変える**"]):
    yy = CONTENT_TOP + Inches(0.3) + Inches(1.5) * i
    label_box(s, MARGIN, yy, Inches(1.0), Inches(1.0), str(i + 1), C_WARN, C_WARN, 32, C_WHITE, True,
              MSO_SHAPE.OVAL)
    textbox(s, MARGIN + Inches(1.3), yy, CONTENT_W - Inches(1.3), Inches(1.0), [txt], size=26,
            anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)

# ---- 確認用:スライド枚数とタイトル一覧
print(f"保存しました:{OUT}")
print(f"スライド枚数:{len(prs.slides)}")
for i, sl in enumerate(prs.slides, 1):
    title = next((sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text),
                 "")
    print(f"{i:2d}. {title}")
